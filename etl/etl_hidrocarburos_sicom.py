#!/usr/bin/env python3
"""
ETL bases SICOM (combustibles líquidos) → PostgreSQL schema `hidrocarburos`.

Fuente: cubos OLAP de SICOM exportados a Excel por la Dirección de Hidrocarburos
y depositados en OneDrive (carpeta "Datos de hidrocarburos").

  Cubos de datos Agentes SICOM 3.xlsx  →  Hoja1        → sicom_agentes
  Cubos 4 SICOM.xlsx                   →  Capacidad    → sicom_capacidad
                                          Compras      → sicom_compras
                                          Compras EDS  → sicom_compras_eds

Por qué este módulo y no el loader genérico `_load_sheets_to_schema`:

  1. Los exports son tablas dinámicas convertidas a tabla y traen una fila
     "Total general" al final que, cargada como dato, duplicaría los agregados.
     Aquí se separa, se usa para CONCILIAR el detalle (metodología §17
     "validación de totales antes y después de la transformación") y se descarta.
  2. "Compras EDS" viene truncado en el límite de filas de Excel (1.048.576):
     hay que detectarlo y dejarlo registrado, porque las cifras de ese tablero
     NO son el total nacional.
  3. Volumen: 1,2 millones de filas. Se leen en streaming (openpyxl read_only)
     y se cargan con COPY, sin materializar DataFrames.

Los cubos son snapshots completos sin clave natural estable (un mismo agente
aparece en varias filas), así que la carga es TRUNCATE + COPY, no UPSERT.

Uso:
    python etl/etl_hidrocarburos_sicom.py --agentes  data/onedrive/SICOM_Agentes.xlsx
    python etl/etl_hidrocarburos_sicom.py --cubos    data/onedrive/SICOM_Cubos4.xlsx
"""

import argparse
import csv
import io
import logging
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import psycopg2

BASE_DIR = Path(__file__).resolve().parent.parent
LOG_DIR = BASE_DIR / "logs"
LOG_DIR.mkdir(exist_ok=True)
SCHEMA_SQL = BASE_DIR / "sql" / "hidrocarburos_sicom_schema.sql"

DEFAULT_XLSX_AGENTES = BASE_DIR / "data" / "onedrive" / "SICOM_Agentes.xlsx"
DEFAULT_XLSX_CUBOS = BASE_DIR / "data" / "onedrive" / "SICOM_Cubos4.xlsx"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - [ETL_SICOM] - %(levelname)s - %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(LOG_DIR / "etl_hidrocarburos_sicom.log", encoding="utf-8"),
    ],
)
logger = logging.getLogger(__name__)

# Límite físico de filas de una hoja de Excel (1.048.576 incluyendo encabezado).
# Un export que llegue exactamente a este número está cortado, no completo.
EXCEL_MAX_DATA_ROWS = 1_048_575

# Marca que Excel pone en la fila de totales de una tabla dinámica.
TOTAL_ROW_LABEL = "total general"

TEXTO, ENTERO, NUMERO = "texto", "entero", "numero"


class HojaSpec:
    """Contrato de una hoja del cubo: encabezados esperados → columnas de la tabla."""

    def __init__(self, hoja, tabla, columnas, medida):
        self.hoja = hoja
        self.tabla = tabla
        self.columnas = columnas  # [(encabezado_excel, columna_pg, tipo), ...]
        self.medida = medida      # columna_pg sobre la que se concilia el total

    @property
    def cols_pg(self):
        return [c[1] for c in self.columnas]

    @property
    def idx_medida(self):
        return self.cols_pg.index(self.medida)


SPEC_AGENTES = HojaSpec(
    hoja="Hoja1",
    tabla="hidrocarburos.sicom_agentes",
    columnas=[
        ("OPERADOR NOMBRE COMERCIAL",  "operador_nombre_comercial",  TEXTO),
        ("TIPO AGENTE",                "tipo_agente",                TEXTO),
        ("SUBTIPO AGENTE",             "subtipo_agente",             TEXTO),
        ("AÑO CREACION",               "anio_creacion",              ENTERO),
        ("BANDERA",                    "bandera",                    TEXTO),
        ("NOMBRE COMERCIAL PROVEEDOR", "nombre_comercial_proveedor", TEXTO),
        ("DEPARTAMENTO INSTALACION",   "departamento_instalacion",   TEXTO),
        ("OPERADA POR",                "operada_por",                TEXTO),
        ("PRODUCTO A COMERCIALIZAR",   "producto_a_comercializar",   TEXTO),
        ("TIPO ESTACION",              "tipo_estacion",              TEXTO),
        ("AGENTE CANTIDAD",            "agente_cantidad",            NUMERO),
    ],
    medida="agente_cantidad",
)

SPEC_CAPACIDAD = HojaSpec(
    hoja="Capacidad",
    tabla="hidrocarburos.sicom_capacidad",
    columnas=[
        ("DEPARTAMENTO INSTALACION", "departamento_instalacion", TEXTO),
        ("MUNICIPIO INSTALACION",    "municipio_instalacion",    TEXTO),
        ("NOMBRE COMERCIAL AGENTE",  "nombre_comercial_agente",  TEXTO),
        ("SUBTIPO AGENTE",           "subtipo_agente",           TEXTO),
        ("TIPO AGENTE",              "tipo_agente",              TEXTO),
        ("PRODUCTO",                 "producto",                 TEXTO),
        ("DESCRIPCION PRODUCTO",     "descripcion_producto",     TEXTO),
        ("CAPACIDAD OPERATIVA",      "capacidad_operativa",      NUMERO),
    ],
    medida="capacidad_operativa",
)

SPEC_COMPRAS = HojaSpec(
    hoja="Compras",
    tabla="hidrocarburos.sicom_compras",
    columnas=[
        ("NOMBRE COMERCIAL AGENTE COMPRADOR", "nombre_comercial_agente_comprador", TEXTO),
        ("TIPO AGENTE COMPRADOR",             "tipo_agente_comprador",             TEXTO),
        ("MUNICIPIO AGENTE PROVEEDOR",        "municipio_agente_proveedor",        TEXTO),
        ("TIPO AGENTE PROVEEDOR",             "tipo_agente_proveedor",             TEXTO),
        ("AÑO CICLO",                         "anio_ciclo",                        ENTERO),
        ("MUNICIPIO PLANTA AGENTE COMPRADOR", "municipio_planta_agente_comprador", TEXTO),
        ("MUNICIPIO PLANTA AGENTE PROVEEDOR", "municipio_planta_agente_proveedor", TEXTO),
        ("VOLUMEN COMPRADO",                  "volumen_comprado",                  NUMERO),
        ("VOLUMEN TRANSITO",                  "volumen_transito",                  NUMERO),
    ],
    medida="volumen_comprado",
)

SPEC_COMPRAS_EDS = HojaSpec(
    hoja="Compras EDS",
    tabla="hidrocarburos.sicom_compras_eds",
    columnas=[
        ("NOMBRE COMERCIAL COMPRADOR", "nombre_comercial_comprador", TEXTO),
        ("NOMBRE COMERCIAL PROVEEDOR", "nombre_comercial_proveedor", TEXTO),
        ("PRODUCTO",                   "producto",                   TEXTO),
        ("AÑO CICLO",                  "anio_ciclo",                 ENTERO),
        ("MES CICLO",                  "mes_ciclo",                  ENTERO),
        ("VOLUMEN COMPRA EDS",         "volumen_compra_eds",         NUMERO),
    ],
    medida="volumen_compra_eds",
)

SPECS_CUBOS = [SPEC_CAPACIDAD, SPEC_COMPRAS, SPEC_COMPRAS_EDS]


# ═══════════════════════════════════════════════════════════════════════════════
# Conexión y DDL
# ═══════════════════════════════════════════════════════════════════════════════

def get_connection():
    """Conexión directa a PostgreSQL (misma config que etl_hidrocarburos.py).

    No usa el pool de `connection_manager` porque su search_path no incluye
    `hidrocarburos` y trae statement_timeout=30s, insuficiente para un COPY
    de 1,2 millones de filas.
    """
    return psycopg2.connect(
        dbname="portal_energetico",
        user="postgres",
        host="localhost",
        port=5432,
        options="-c search_path=hidrocarburos,public",
    )


def ensure_schema(conn) -> None:
    """Crea schema/tablas/índices si no existen ejecutando el DDL idempotente."""
    if not SCHEMA_SQL.exists():
        raise RuntimeError(f"DDL no encontrado: {SCHEMA_SQL}")
    with conn.cursor() as cur:
        cur.execute(SCHEMA_SQL.read_text(encoding="utf-8"))
    conn.commit()
    logger.info("✅ Schema hidrocarburos (SICOM) verificado/creado")


# ═══════════════════════════════════════════════════════════════════════════════
# Normalización de celdas
# ═══════════════════════════════════════════════════════════════════════════════

def _texto(v):
    """Normaliza una celda de dimensión. Devuelve None para vacíos.

    El guion '-' NO se convierte a NULL: es un valor que el propio cubo emite
    (p.ej. BANDERA sin bandera) y convertirlo sería reinterpretar la fuente.
    Los tableros lo filtran explícitamente donde corresponde.
    """
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "null"):
        return None
    return s


def _numero(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)  # v != v  → NaN
    s = str(v).strip().replace(",", "")
    if s == "" or s.lower() in ("nan", "none", "null", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _entero(v):
    f = _numero(v)
    return None if f is None else int(f)


_CONVERSORES = {TEXTO: _texto, ENTERO: _entero, NUMERO: _numero}


# ═══════════════════════════════════════════════════════════════════════════════
# Lectura en streaming + carga
# ═══════════════════════════════════════════════════════════════════════════════

def _validar_encabezados(fila, spec: HojaSpec) -> None:
    """Falla ruidosamente si el cubo cambió de estructura.

    Es la defensa contra el riesgo que la propia metodología §19 señala: que una
    actualización de SICOM cambie las columnas y el tablero muestre otra cosa
    sin que nadie lo note.
    """
    leidos = [(_texto(c) or "").upper() for c in fila[: len(spec.columnas)]]
    esperados = [c[0].upper() for c in spec.columnas]
    if leidos != esperados:
        faltan = [e for e in esperados if e not in leidos]
        sobran = [l for l in leidos if l and l not in esperados]
        raise RuntimeError(
            f"[{spec.hoja}] Encabezados inesperados.\n"
            f"  esperados: {esperados}\n"
            f"  leídos   : {leidos}\n"
            f"  faltan   : {faltan}\n"
            f"  sobran   : {sobran}"
        )


def _procesar_hoja(ws, spec: HojaSpec):
    """Lee la hoja en streaming → (buffer CSV, estadísticas de conciliación)."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")

    n_cols = len(spec.columnas)
    conversores = [_CONVERSORES[c[2]] for c in spec.columnas]
    idx_medida = spec.idx_medida

    filas = 0
    suma_detalle = 0.0
    total_declarado = None
    primer_clave = None
    ultimo_clave = None
    filas_brutas = 0

    for i, fila in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            _validar_encabezados(fila, spec)
            continue

        fila = fila[:n_cols]
        if all(c is None for c in fila):
            continue
        filas_brutas += 1

        clave = _texto(fila[0])

        # Fila de totales de la tabla dinámica: se usa para conciliar, no se carga.
        if clave and clave.lower() == TOTAL_ROW_LABEL:
            total_declarado = _numero(fila[idx_medida])
            continue

        valores = [conv(v) for conv, v in zip(conversores, fila)]

        medida = valores[idx_medida]
        if medida is not None:
            suma_detalle += medida

        if primer_clave is None:
            primer_clave = clave
        ultimo_clave = clave

        writer.writerow(["" if v is None else v for v in valores])
        filas += 1

    buf.seek(0)

    truncado = filas_brutas >= EXCEL_MAX_DATA_ROWS
    if total_declarado is not None and total_declarado != 0:
        gap_abs = total_declarado - suma_detalle
        gap_pct = gap_abs / total_declarado * 100.0
        concilia = abs(gap_pct) < 0.5  # tolerancia de redondeo del cubo
    else:
        gap_abs = gap_pct = None
        concilia = None

    stats = {
        "filas": filas,
        "medida": spec.medida,
        "suma_detalle": suma_detalle,
        "total_declarado": total_declarado,
        "gap_abs": gap_abs,
        "gap_pct": gap_pct,
        "concilia": concilia,
        "truncado": truncado,
        "primer_clave": primer_clave,
        "ultimo_clave": ultimo_clave,
    }
    return buf, stats


def _cargar(conn, spec: HojaSpec, buf, stats) -> None:
    """TRUNCATE + COPY de la tabla destino (dentro de la transacción del caller)."""
    cols = ", ".join(spec.cols_pg)
    with conn.cursor() as cur:
        cur.execute(f"TRUNCATE TABLE {spec.tabla} RESTART IDENTITY")
        cur.copy_expert(
            f"COPY {spec.tabla} ({cols}) FROM STDIN WITH (FORMAT csv, NULL '')",
            buf,
        )
        cur.execute(f"SELECT COUNT(*) FROM {spec.tabla}")
        n_real = cur.fetchone()[0]

    if n_real != stats["filas"]:
        raise RuntimeError(
            f"[{spec.tabla}] COPY cargó {n_real} filas, se esperaban {stats['filas']}"
        )


def _registrar_control(conn, spec: HojaSpec, archivo: str, stats, fecha_dato) -> None:
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO hidrocarburos.sicom_carga_control
                (tabla, archivo_fuente, hoja, filas_cargadas, medida,
                 suma_detalle, total_general_declarado, gap_abs, gap_pct,
                 concilia, truncado_excel, primer_valor_clave, ultimo_valor_clave,
                 fecha_dato)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (
                spec.tabla.split(".")[-1], archivo, spec.hoja, stats["filas"],
                stats["medida"], stats["suma_detalle"], stats["total_declarado"],
                stats["gap_abs"], stats["gap_pct"], stats["concilia"],
                stats["truncado"], stats["primer_clave"], stats["ultimo_clave"],
                fecha_dato,
            ),
        )


def _log_stats(spec: HojaSpec, stats) -> None:
    logger.info("  %s → %s filas", spec.hoja, f"{stats['filas']:,}")
    if stats["truncado"]:
        logger.warning(
            "  ⚠️  [%s] TRUNCADO en el límite de Excel (%s filas). "
            "Cobertura parcial: de %r a %r. Las cifras NO son el total nacional.",
            spec.hoja, f"{EXCEL_MAX_DATA_ROWS:,}",
            stats["primer_clave"], stats["ultimo_clave"],
        )
    if stats["concilia"] is None:
        logger.info("  · sin fila 'Total general' — no hay contra qué conciliar")
    elif stats["concilia"]:
        logger.info(
            "  ✅ concilia con el total del cubo (%s, desvío %.4f%%)",
            f"{stats['total_declarado']:,.0f}", stats["gap_pct"],
        )
    else:
        logger.warning(
            "  ⚠️  [%s] NO CONCILIA: detalle %s vs total declarado %s "
            "(faltan %s = %.2f%% de %s)",
            spec.hoja, f"{stats['suma_detalle']:,.0f}",
            f"{stats['total_declarado']:,.0f}", f"{stats['gap_abs']:,.0f}",
            stats["gap_pct"], stats["medida"],
        )


def _procesar_archivo(xlsx_path: Path, specs, fecha_fuente=None) -> dict:
    """Carga todas las hojas de un archivo en UNA transacción atómica."""
    if not xlsx_path.exists():
        raise FileNotFoundError(f"No existe: {xlsx_path}")

    logger.info("📂 %s (%.1f MB)", xlsx_path.name, xlsx_path.stat().st_size / 1024 / 1024)
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        pendientes = []
        for spec in specs:
            if spec.hoja not in wb.sheetnames:
                raise RuntimeError(
                    f"Hoja {spec.hoja!r} no está en {xlsx_path.name}. "
                    f"Hojas presentes: {wb.sheetnames}"
                )
            buf, stats = _procesar_hoja(wb[spec.hoja], spec)
            _log_stats(spec, stats)
            pendientes.append((spec, buf, stats))
    finally:
        wb.close()

    conn = get_connection()
    try:
        ensure_schema(conn)
        for spec, buf, stats in pendientes:
            _cargar(conn, spec, buf, stats)
            _registrar_control(conn, spec, xlsx_path.name, stats, fecha_fuente)
        conn.commit()
        logger.info("✅ %s: transacción commiteada", xlsx_path.name)
    except Exception:
        conn.rollback()
        logger.error("❌ %s — ROLLBACK ejecutado", xlsx_path.name, exc_info=True)
        raise
    finally:
        conn.close()

    return {spec.tabla.split(".")[-1]: stats for spec, _, stats in pendientes}


# ═══════════════════════════════════════════════════════════════════════════════
# Puntos de entrada (los usa etl_sharepoint_sync)
# ═══════════════════════════════════════════════════════════════════════════════

def _con_linaje(paso: str, fn):
    """Envuelve la corrida en ontologia.etl_lineage (nunca tumba el ETL)."""
    try:
        sys.path.insert(0, str(BASE_DIR))
        from infrastructure.observability.lineage import registrar_paso
    except Exception:
        return fn()

    with registrar_paso("sicom_hidrocarburos", paso) as reg:
        res = fn()
        reg.filas(sum(s["filas"] for s in res.values()))
    return res


def run_etl_agentes(xlsx_path: Path, fecha_fuente: datetime | None = None) -> dict:
    return _con_linaje(
        "sicom_agentes",
        lambda: _procesar_archivo(Path(xlsx_path), [SPEC_AGENTES], fecha_fuente),
    )


def run_etl_cubos(xlsx_path: Path, fecha_fuente: datetime | None = None) -> dict:
    return _con_linaje(
        "sicom_cubos4",
        lambda: _procesar_archivo(Path(xlsx_path), SPECS_CUBOS, fecha_fuente),
    )


def main():
    parser = argparse.ArgumentParser(description="ETL cubos SICOM → PostgreSQL")
    parser.add_argument("--agentes", type=Path, nargs="?", const=DEFAULT_XLSX_AGENTES,
                        help="Excel del cubo de Agentes")
    parser.add_argument("--cubos", type=Path, nargs="?", const=DEFAULT_XLSX_CUBOS,
                        help="Excel de Cubos 4 (Capacidad/Compras/Compras EDS)")
    args = parser.parse_args()

    if not args.agentes and not args.cubos:
        parser.error("Indica al menos --agentes o --cubos")

    try:
        if args.agentes:
            run_etl_agentes(args.agentes)
        if args.cubos:
            run_etl_cubos(args.cubos)
    except Exception as e:
        logger.error("ETL SICOM falló: %s", e)
        sys.exit(1)
    logger.info("🏁 ETL SICOM completado")


if __name__ == "__main__":
    main()
