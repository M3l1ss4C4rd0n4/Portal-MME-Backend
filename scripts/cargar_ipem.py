#!/usr/bin/env python3
"""
Cargar el IPEM (Índice de Pobreza Energética Multidimensional) desde el Excel del
área a PostgreSQL.

Dato FIJO de entrega manual: no hay ETL ni cron. Se corre una vez al año, cuando el
área entregue el Excel del año siguiente. El script es idempotente — vuelve a dejar
las tablas con exactamente el contenido del Excel, así que correrlo dos veces no
duplica nada.

NO notifica por ningún canal: no manda correo ni Telegram. Solo escribe en la BD y
reporta por salida estándar.

Requiere las tablas de sql/migrations/047_ipem_2024_2025.sql ya aplicadas, y un
usuario con INSERT (mme_user es de solo lectura; el backend y los ETL de este repo
se conectan como postgres).

Uso:
    /home/admonctrlxm/server/venv/bin/python3 /home/admonctrlxm/server/scripts/cargar_ipem.py
    # solo validar el Excel, sin tocar la base:
    /home/admonctrlxm/server/venv/bin/python3 /home/admonctrlxm/server/scripts/cargar_ipem.py --dry-run
"""

import argparse
import os
import re
import sys
import unicodedata

import openpyxl
import psycopg2

EXCEL_PATH = "/home/admonctrlxm/server/data/IPEM/IPEM.xlsx"

HOJA_MUNICIPAL = "IPEM Municip"
HOJA_DEPARTAMENTAL = "IPEM Departamental"

# Totales que debe traer la entrega. Si no cuadran, el Excel cambió de forma y hay que
# revisarlo con el área antes de publicar: fallar es preferible a cargar a medias.
MUNICIPIOS_ESPERADOS = 1121
DEPARTAMENTOS_ESPERADOS = 33

# Angelópolis (05036) llegó sin valor de 2024 en la entrega original, aunque sí trae
# etiqueta de quintil. Es el único hueco conocido; si aparecen más, se avisa.
NULOS_2024_CONOCIDOS = 1

# Vocabulario cerrado de la columna de quintil. Si la entrega del año próximo escribe
# «Quintil 1» o «Muy bajo» en vez de estas cinco, el tablero no mostraría ninguna
# etiqueta y la gráfica de distribución se iría a cero SIN que nadie viera un error:
# el frontend convierte lo que no reconoce en «—». Por eso se valida aquí y se aborta.
QUINTILES_VALIDOS = frozenset({"MUY BAJO", "BAJO", "MEDIO", "ALTO", "MUY ALTO"})

# Orden de menor a mayor incidencia, para comprobar que las etiquetas sean coherentes
# con los valores (el rango de una etiqueta no debe solaparse con el de la siguiente).
QUINTILES_ORDEN = ("MUY BAJO", "BAJO", "MEDIO", "ALTO", "MUY ALTO")

DB_CONFIG = dict(
    host=os.getenv("PGHOST", "localhost"),
    port=os.getenv("PGPORT", "5432"),
    database=os.getenv("PGDATABASE", "portal_energetico"),
    user=os.getenv("PGUSER", "postgres"),
)
# Solo se pasa la contraseña si viene en el entorno: con la clave vacía explícita,
# libpq no consulta ~/.pgpass y la autenticación falla en esta máquina.
if os.getenv("PGPASSWORD"):
    DB_CONFIG["password"] = os.environ["PGPASSWORD"]


class ErrorDeFuente(Exception):
    """El Excel no tiene la forma esperada. Se aborta sin tocar la base."""


def normalizar(texto) -> str:
    """Clave de comparación de nombres: sin tildes, sin puntuación, en mayúsculas.

    El Excel escribe «BOGOTÁ, D.C.» y «ARCHIPIELAGO DE SAN ANDRÉS, PROVIDENCIA Y
    SANTA CATALINA»; la hoja municipal usa «Antioquia» en minúsculas de título. Sin
    normalizar, el cruce entre las dos hojas falla en esos nombres.
    """
    s = unicodedata.normalize("NFD", str(texto))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return " ".join(s.split())


def texto(valor) -> str | None:
    if valor is None:
        return None
    s = str(valor).strip()
    return s or None


def indice(valor) -> float | None:
    """Valor del IPEM como fracción 0-1, o None si la celda viene vacía."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        n = float(valor)
    except (TypeError, ValueError) as exc:
        raise ErrorDeFuente(f"valor de IPEM no numérico: {valor!r}") from exc
    if not 0 <= n <= 1:
        raise ErrorDeFuente(
            f"valor de IPEM fuera de 0-1: {n}. El Excel entrega fracciones "
            f"(0,1917 = 19,17 %); si ahora llega en porcentaje, hay que dividir por 100."
        )
    return n


def leer_municipios(ws) -> list[dict]:
    """Hoja municipal → filas listas para insertar.

    Columnas: 0 Departamento · 1 Municipio · 2 Cod Municipio · 3 IPEM 2024 ·
    4 IPEM 2025 · 5 QUINTIL 2024 · 6 QUINTIL 2025.
    """
    filas = []
    for cruda in list(ws.iter_rows(values_only=True))[1:]:
        if cruda[0] is None:
            continue
        # Los códigos llegan sin relleno (5001 en vez de 05001): DIVIPOLA son 5
        # dígitos y los 2 primeros son el departamento, que es la clave con la que
        # el tablero cruza contra el geojson.
        codigo = str(cruda[2]).strip().split(".")[0].zfill(5)
        if not codigo.isdigit() or len(codigo) != 5:
            raise ErrorDeFuente(f"código de municipio inesperado: {cruda[2]!r}")
        filas.append(
            {
                "cod_municipio": codigo,
                "cod_departamento": codigo[:2],
                "departamento": texto(cruda[0]),
                "municipio": texto(cruda[1]),
                "ipem_2024": indice(cruda[3]),
                "ipem_2025": indice(cruda[4]),
                "quintil_2024": texto(cruda[5]),
                "quintil_2025": texto(cruda[6]),
            }
        )
    return filas


def leer_departamentos(
    ws,
    codigo_por_nombre: dict[str, str],
    nombre_canonico: dict[str, str],
) -> tuple[list[dict], list[str]]:
    """Hoja departamental → filas listas para insertar, más los avisos de grafía.

    Columnas: 0 REGIÓN · 1 Nombre Departamento · 2 IPEM 2024 · 3 IPEM 2025.

    La hoja no trae código, así que se toma del mapa construido con los códigos
    municipales. Se hace así, y no por nombre contra el geojson, porque el geojson
    del portal escribe Bogotá como «SANTAFE DE BOGOTA D.C» y ningún normalizador
    razonable lo iguala a «BOGOTÁ, D.C.».

    El nombre que se GUARDA es el canónico de la hoja municipal, no el de esta hoja.
    Las dos hojas escriben el mismo departamento distinto —«Antioquia» en la municipal
    y «ANTIOQUIA» en la departamental, en los 33 casos— y conservar cada grafía hacía
    que el tablero mostrara el mismo departamento en mayúsculas en el mapa y en tipo
    título en la tabla. Una sola fuente de verdad para el nombre visible.
    """
    filas = []
    avisos = []
    for cruda in list(ws.iter_rows(values_only=True))[1:]:
        if cruda[0] is None:
            continue
        nombre_hoja = texto(cruda[1])
        clave = normalizar(nombre_hoja)
        codigo = codigo_por_nombre.get(clave)
        if codigo is None:
            raise ErrorDeFuente(
                f"el departamento «{nombre_hoja}» de la hoja departamental no aparece en "
                f"la hoja municipal, así que no se le puede asignar código DIVIPOLA"
            )
        canonico = nombre_canonico[clave]
        if canonico != nombre_hoja:
            avisos.append(f"{codigo}: «{nombre_hoja}» → se guarda «{canonico}»")
        filas.append(
            {
                "cod_departamento": codigo,
                "departamento": canonico,
                "region": texto(cruda[0]),
                "ipem_2024": indice(cruda[2]),
                "ipem_2025": indice(cruda[3]),
            }
        )
    return filas, avisos


def validar(municipios: list[dict], departamentos: list[dict]) -> list[str]:
    """Comprueba la entrega antes de escribir. Devuelve avisos; lanza si es grave."""
    if len(municipios) != MUNICIPIOS_ESPERADOS:
        raise ErrorDeFuente(
            f"se esperaban {MUNICIPIOS_ESPERADOS} municipios y llegaron {len(municipios)}"
        )
    if len(departamentos) != DEPARTAMENTOS_ESPERADOS:
        raise ErrorDeFuente(
            f"se esperaban {DEPARTAMENTOS_ESPERADOS} departamentos y llegaron "
            f"{len(departamentos)}"
        )

    codigos = [m["cod_municipio"] for m in municipios]
    if len(set(codigos)) != len(codigos):
        repetidos = sorted({c for c in codigos if codigos.count(c) > 1})
        raise ErrorDeFuente(f"códigos de municipio repetidos: {repetidos}")

    codigos_depto = {d["cod_departamento"] for d in departamentos}
    huerfanos = sorted({m["cod_departamento"] for m in municipios} - codigos_depto)
    if huerfanos:
        raise ErrorDeFuente(
            f"hay municipios cuyo departamento no está en la hoja departamental: {huerfanos}"
        )

    # Vocabulario cerrado de las etiquetas. Esto es lo que impide que una entrega con
    # otra nomenclatura entre sin ruido y deje el tablero sin clasificación visible.
    for anio in ("2024", "2025"):
        vistas = {m[f"quintil_{anio}"] for m in municipios if m[f"quintil_{anio}"]}
        desconocidas = sorted(vistas - QUINTILES_VALIDOS)
        if desconocidas:
            raise ErrorDeFuente(
                f"la columna QUINTIL {anio} trae etiquetas que el tablero no sabe "
                f"interpretar: {desconocidas}. Se esperaba exactamente "
                f"{sorted(QUINTILES_VALIDOS)}. El tablero las mostraría como «—» en "
                f"todas las filas sin avisar, así que se aborta la carga."
            )

    avisos = []

    # Coherencia entre la etiqueta y el valor: el rango de cada quintil no debe
    # solaparse con el del siguiente. Un solapamiento significaría que las etiquetas
    # no corresponden a los valores de su propia fila (columnas desalineadas, por
    # ejemplo), y eso no lo detecta ningún conteo.
    for anio in ("2024", "2025"):
        rangos: dict[str, tuple[float, float]] = {}
        for etiqueta in QUINTILES_ORDEN:
            vals = [
                m[f"ipem_{anio}"]
                for m in municipios
                if m[f"quintil_{anio}"] == etiqueta and m[f"ipem_{anio}"] is not None
            ]
            if vals:
                rangos[etiqueta] = (min(vals), max(vals))
        presentes = [q for q in QUINTILES_ORDEN if q in rangos]
        for menor, mayor in zip(presentes, presentes[1:]):
            if rangos[menor][1] >= rangos[mayor][0]:
                avisos.append(
                    f"QUINTIL {anio}: el rango de «{menor}» se solapa con el de "
                    f"«{mayor}» ({rangos[menor][1]:.6f} ≥ {rangos[mayor][0]:.6f}) — "
                    f"revisar si las etiquetas corresponden a los valores de su fila"
                )

    nulos_2024 = [m["cod_municipio"] for m in municipios if m["ipem_2024"] is None]
    nulos_2025 = [m["cod_municipio"] for m in municipios if m["ipem_2025"] is None]
    if len(nulos_2024) != NULOS_2024_CONOCIDOS:
        avisos.append(
            f"municipios sin IPEM 2024: {len(nulos_2024)} (se esperaba "
            f"{NULOS_2024_CONOCIDOS}, Angelópolis) → {nulos_2024}"
        )
    if nulos_2025:
        avisos.append(f"municipios sin IPEM 2025: {len(nulos_2025)} → {nulos_2025}")

    # Un quintil sin valor es justo la inconsistencia que trae Angelópolis: se deja
    # pasar, pero se nombra, porque decide cómo se pinta esa fila en el tablero.
    for m in municipios:
        if m["ipem_2024"] is None and m["quintil_2024"]:
            avisos.append(
                f"{m['municipio']} ({m['cod_municipio']}) trae quintil 2024 "
                f"«{m['quintil_2024']}» pero no trae valor de IPEM 2024"
            )
    return avisos


def resumen(municipios: list[dict], departamentos: list[dict]) -> None:
    """Promedios y conteo por quintil, para contrastar contra el tablero."""
    for etiqueta, filas in (("municipal", municipios), ("departamental", departamentos)):
        for anio in ("2024", "2025"):
            vals = [f[f"ipem_{anio}"] for f in filas if f[f"ipem_{anio}"] is not None]
            print(
                f"  {etiqueta:<14} {anio}: n={len(vals):>4} "
                f"min={min(vals) * 100:5.2f} %  max={max(vals) * 100:5.2f} %  "
                f"promedio={sum(vals) / len(vals) * 100:5.2f} %"
            )
    for anio in ("2024", "2025"):
        conteo: dict[str, int] = {}
        for m in municipios:
            q = m[f"quintil_{anio}"] or "sin etiqueta"
            conteo[q] = conteo.get(q, 0) + 1
        partes = [f"{q}={conteo[q]}" for q in QUINTILES_ORDEN if q in conteo]
        extra = [f"{q}={n}" for q, n in conteo.items() if q not in QUINTILES_ORDEN]
        print(f"  quintiles municipales {anio}: " + "  ".join(partes + extra))


def perfilar(municipios: list[dict], departamentos: list[dict]) -> None:
    """Perfilado de la entrega: lo que conviene mirar aunque no sea un error.

    Ninguna de estas cifras aborta la carga. Están para que el cambio de un año a
    otro salte a la vista: si los homónimos o las áreas no municipalizadas cambian de
    número, la fuente cambió de criterio y hay que preguntar antes de publicar.
    """
    nombres: dict[str, int] = {}
    for m in municipios:
        clave = normalizar(m["municipio"])
        nombres[clave] = nombres.get(clave, 0) + 1
    homonimos = {k: n for k, n in nombres.items() if n > 1}

    # Las «(CD)» son corregimientos departamentales / áreas no municipalizadas: no son
    # municipios, aunque viajen en la misma hoja. Importa para no llamar «municipios»
    # al total en el tablero ni en el informe.
    no_municipalizadas = [m for m in municipios if re.search(r"\(\s*(CD|ANM|CP)\s*\)", m["municipio"])]

    print(
        f"  perfilado: {len(homonimos)} nombres de municipio repetidos entre "
        f"departamentos · {len(no_municipalizadas)} áreas no municipalizadas «(CD)» · "
        f"{len(municipios) - len(no_municipalizadas)} municipios propiamente dichos"
    )

    # Atípicos por el criterio de Tukey. Se reportan, no se eliminan: en este índice
    # los extremos son municipios de la Amazonía-Orinoquía con privación realmente
    # alta, no errores de captura. Recortarlos borraría justo lo que el informe busca.
    for anio in ("2024", "2025"):
        vals = sorted(m[f"ipem_{anio}"] for m in municipios if m[f"ipem_{anio}"] is not None)
        if len(vals) < 4:
            continue
        q1 = vals[len(vals) // 4]
        q3 = vals[(3 * len(vals)) // 4]
        iqr = q3 - q1
        fuera = [v for v in vals if v < q1 - 1.5 * iqr or v > q3 + 1.5 * iqr]
        print(
            f"  atípicos {anio} (Tukey, solo informativo): {len(fuera)} de {len(vals)} "
            f"· máximo {max(vals) * 100:.2f} %"
        )

    sin_region = [d["cod_departamento"] for d in departamentos if not d["region"]]
    if sin_region:
        print(f"  AVISO: departamentos sin región: {sin_region}")


def escribir(conn, municipios: list[dict], departamentos: list[dict]) -> None:
    """Reemplaza el contenido de las dos tablas en una sola transacción.

    Las dos tablas se truncan en UNA sentencia y los departamentos se insertan
    primero: desde la migración 048 hay una llave foránea de ipem_municipal hacia
    ipem_departamental, así que truncar o insertar en el otro orden la violaría.
    """
    with conn.cursor() as cur:
        cur.execute(
            "TRUNCATE sector_energetico.ipem_municipal, "
            "sector_energetico.ipem_departamental"
        )
        cur.executemany(
            """
            INSERT INTO sector_energetico.ipem_departamental
                (cod_departamento, departamento, region, ipem_2024, ipem_2025)
            VALUES (%(cod_departamento)s, %(departamento)s, %(region)s,
                    %(ipem_2024)s, %(ipem_2025)s)
            """,
            departamentos,
        )
        cur.executemany(
            """
            INSERT INTO sector_energetico.ipem_municipal
                (cod_municipio, cod_departamento, departamento, municipio,
                 ipem_2024, ipem_2025, quintil_2024, quintil_2025)
            VALUES (%(cod_municipio)s, %(cod_departamento)s, %(departamento)s,
                    %(municipio)s, %(ipem_2024)s, %(ipem_2025)s,
                    %(quintil_2024)s, %(quintil_2025)s)
            """,
            municipios,
        )
    conn.commit()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="validar y resumir el Excel sin escribir en la base",
    )
    parser.add_argument("--excel", default=EXCEL_PATH, help="ruta del Excel a cargar")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"ERROR: no existe el Excel {args.excel}", file=sys.stderr)
        return 1

    print(f"Leyendo {args.excel}")
    wb = openpyxl.load_workbook(args.excel, read_only=True, data_only=True)
    faltantes = [h for h in (HOJA_MUNICIPAL, HOJA_DEPARTAMENTAL) if h not in wb.sheetnames]
    if faltantes:
        print(
            f"ERROR: el Excel no trae las hojas {faltantes}. Tiene: {wb.sheetnames}",
            file=sys.stderr,
        )
        return 1

    try:
        municipios = leer_municipios(wb[HOJA_MUNICIPAL])
        codigo_por_nombre = {
            normalizar(m["departamento"]): m["cod_departamento"] for m in municipios
        }
        # La hoja municipal es la fuente de verdad del nombre visible: escribe
        # «Antioquia» y no «ANTIOQUIA», que es lo que el tablero quiere mostrar.
        nombre_canonico = {
            normalizar(m["departamento"]): m["departamento"] for m in municipios
        }
        departamentos, avisos_grafia = leer_departamentos(
            wb[HOJA_DEPARTAMENTAL], codigo_por_nombre, nombre_canonico
        )
        avisos = validar(municipios, departamentos)
    except ErrorDeFuente as exc:
        print(f"ERROR en la fuente: {exc}", file=sys.stderr)
        print("No se escribió nada en la base.", file=sys.stderr)
        return 1

    print(f"OK: {len(municipios)} municipios y {len(departamentos)} departamentos")
    resumen(municipios, departamentos)
    perfilar(municipios, departamentos)
    if avisos_grafia:
        print(
            f"  grafía unificada: {len(avisos_grafia)} de {len(departamentos)} nombres de "
            f"la hoja departamental se guardan con la grafía de la hoja municipal"
        )
    for aviso in avisos:
        print(f"  AVISO: {aviso}")

    if args.dry_run:
        print("--dry-run: no se escribió en la base.")
        return 0

    try:
        conn = psycopg2.connect(**DB_CONFIG)
    except psycopg2.Error as exc:
        print(f"ERROR conectando a PostgreSQL: {exc}", file=sys.stderr)
        print(
            "Las tablas se escriben con un usuario que tenga INSERT; mme_user es de "
            "solo lectura. Exportar PGUSER/PGPASSWORD del rol adecuado.",
            file=sys.stderr,
        )
        return 1

    try:
        escribir(conn, municipios, departamentos)
    except psycopg2.Error as exc:
        conn.rollback()
        print(f"ERROR escribiendo en PostgreSQL: {exc}", file=sys.stderr)
        print(
            "Transacción revertida; las tablas quedaron como estaban. Si la queja es "
            "que no existen, falta aplicar sql/migrations/047_ipem_2024_2025.sql.",
            file=sys.stderr,
        )
        return 1
    finally:
        conn.close()

    print("Cargado en sector_energetico.ipem_municipal e ipem_departamental.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
